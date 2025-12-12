"""
Generate PowerPoint presentation from the final report
"""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import json

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

def add_title_slide(title, subtitle):
    """Add title slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    return slide

def add_section_slide(title):
    """Add section divider slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    slide.shapes.title.text = title
    return slide

def add_content_slide(title, bullet_points):
    """Add slide with title and bullet points"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    content = slide.placeholders[1].text_frame
    content.clear()

    for point in bullet_points:
        p = content.add_paragraph()
        p.text = point
        p.level = 0
        p.font.size = Pt(18)

    return slide

def add_two_column_slide(title, left_content, right_content):
    """Add slide with two columns"""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title

    # Left column
    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.5), Inches(5))
    tf = left_box.text_frame
    for item in left_content:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)

    # Right column
    right_box = slide.shapes.add_textbox(Inches(5.2), Inches(1.5), Inches(4.5), Inches(5))
    tf = right_box.text_frame
    for item in right_content:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)

    return slide

# Slide 1: Title
add_title_slide(
    "Airline Assistant",
    "Knowledge Graph-Based Flight Recommendation System\nFinal Presentation"
)

# Slide 2: Agenda
add_content_slide(
    "Agenda",
    [
        "1. System Architecture",
        "2. Retrieval Strategy & Examples",
        "3. LLM Model Comparison",
        "4. Error Analysis",
        "5. Improvements & Enhancements",
        "6. Remaining Limitations",
        "7. Conclusion & Future Work"
    ]
)

# Section 1: System Architecture
add_section_slide("1. System Architecture")

add_content_slide(
    "System Overview",
    [
        "Multi-stage RAG pipeline combining:",
        "  • Knowledge Graph (Neo4j) - 2,000 flight journeys",
        "  • NLP Preprocessing (SpaCy) - Intent & entity extraction",
        "  • Hybrid Retrieval - Symbolic (Cypher) + Semantic (embeddings)",
        "  • LLM Generation (OpenRouter) - 3 models tested",
        "  • Web UI (Streamlit) - Interactive interface"
    ]
)

add_content_slide(
    "System Flow",
    [
        "User Query → Preprocessing → Retrieval → LLM → Response",
        "",
        "1. Preprocessing: Extract intent & entities",
        "2. Baseline Retrieval: Cypher queries (exact matching)",
        "3. Embedding Retrieval: TF-IDF/BoW (semantic search)",
        "4. Context Combination: Merge, deduplicate, rank",
        "5. Prompt Construction: Structured 3-part prompt",
        "6. LLM Generation: Generate natural language answer"
    ]
)

add_content_slide(
    "Knowledge Graph Schema",
    [
        "Nodes:",
        "  • Journey: class, delay, food_score, miles, dates",
        "  • Flight: flight_number, departure/arrival times",
        "  • Airport: station_code, city, country",
        "",
        "Relationships:",
        "  • Journey -[:ON]-> Flight",
        "  • Flight -[:DEPARTS_FROM]-> Airport",
        "  • Flight -[:ARRIVES_AT]-> Airport"
    ]
)

# Section 2: Retrieval Strategy
add_section_slide("2. Retrieval Strategy")

add_two_column_slide(
    "Retrieval Approaches",
    [
        "Baseline (Symbolic):",
        "• Cypher queries",
        "• Exact attribute matching",
        "• Fast & precise",
        "• Example: origin='JFK'",
        "",
        "Strengths:",
        "• No false positives",
        "• Handles constraints well",
        "• Fast execution"
    ],
    [
        "Embeddings (Semantic):",
        "• TF-IDF / Bag-of-Words",
        "• Cosine similarity",
        "• Handles natural language",
        "• Example: 'comfortable flight'",
        "",
        "Strengths:",
        "• Semantic understanding",
        "• Works with variations",
        "• No exact match needed"
    ]
)

add_content_slide(
    "Example Query: Business Class to LAX",
    [
        "Query: 'Find business class flight from JFK to LAX with low delay'",
        "",
        "Baseline Results (5 flights):",
        "  • Exact matches on: JFK, LAX, Business class",
        "  • Filtered by: delay < 30 minutes",
        "",
        "Embedding Results (10 flights):",
        "  • Semantically similar journeys",
        "  • Cosine similarity > 0.7",
        "",
        "Combined: 14 unique flights, ranked by delay → food → similarity"
    ]
)

# Section 3: LLM Comparison
add_section_slide("3. LLM Model Comparison")

# Load results
with open('model_comparison_results.json', 'r') as f:
    results_data = json.load(f)

add_content_slide(
    "Models Tested",
    [
        "1. Mistral 7B Instruct",
        "   • mistralai/mistral-7b-instruct",
        "",
        "2. Meta Llama 3.2 3B",
        "   • meta-llama/llama-3.2-3b-instruct",
        "",
        "3. Google Gemma 2 9B",
        "   • google/gemma-2-9b-it"
    ]
)

add_two_column_slide(
    "Quantitative Results",
    [
        "Avg Response Time:",
        "  • Mistral 7B: 7.19s",
        "  • Llama 3.2: 8.25s",
        "  • Gemma 2: 3.87s ⭐",
        "",
        "Avg Tokens Used:",
        "  • Mistral 7B: 1,050",
        "  • Llama 3.2: 1,124",
        "  • Gemma 2: 986 ⭐"
    ],
    [
        "Overall Quality (1-5):",
        "  • Mistral 7B: 3.61",
        "  • Llama 3.2: 4.65 ⭐",
        "  • Gemma 2: 4.05",
        "",
        "Success Rate:",
        "  • All models: 100%",
        "",
        "Cost per Request:",
        "  • Gemma 2 most economical"
    ]
)

add_content_slide(
    "Qualitative Analysis",
    [
        "Llama 3.2 3B - Highest Quality:",
        "  • Perfect accuracy (5.0/5) - no hallucination",
        "  • Most complete responses (5.0/5)",
        "  • Best for detailed explanations",
        "",
        "Gemma 2 9B - Best Performance:",
        "  • 2x faster than competitors",
        "  • Most cost-effective",
        "  • Best for production/high traffic",
        "",
        "Mistral 7B - Needs Improvement:",
        "  • Inconsistent (one test returned 1 word)",
        "  • Low relevance scores"
    ]
)

add_content_slide(
    "Recommendation",
    [
        "🏆 Production Deployment: Gemma 2 9B",
        "  • Best cost-performance balance",
        "  • Fastest response time (better UX)",
        "  • Good quality for most queries",
        "",
        "Alternative: Llama 3.2 3B for Premium",
        "  • When detailed explanations needed",
        "  • When accuracy is critical",
        "  • When speed less important",
        "",
        "Not Recommended: Mistral 7B",
        "  • Unreliable response quality"
    ]
)

# Section 4: Error Analysis
add_section_slide("4. Error Analysis")

add_content_slide(
    "Preprocessing Errors",
    [
        "1. Ambiguous Intent:",
        "   • Problem: 'Tell me about flights' - too vague",
        "   • Solution: Add clarification prompts",
        "",
        "2. Entity Extraction Failures:",
        "   • Problem: 'New York' not mapped to JFK/LGA/EWR",
        "   • Solution: City-to-airport alias mapping",
        "",
        "3. Date Parsing:",
        "   • Problem: 'next Tuesday' not recognized",
        "   • Solution: Relative date parsing library"
    ]
)

add_content_slide(
    "Retrieval Errors",
    [
        "1. No Baseline Results:",
        "   • Problem: 'cheap flights' - no price data in KG",
        "   • Solution: Add pricing to schema or inform user",
        "",
        "2. Irrelevant Embedding Results:",
        "   • Problem: Query mentions Wi-Fi, but not in KG",
        "   • Solution: Filter unavailable features",
        "",
        "3. Ranking Issues:",
        "   • Problem: 'Best' is subjective",
        "   • Solution: Learn preferences or ask clarifying questions"
    ]
)

add_content_slide(
    "LLM Generation Errors",
    [
        "1. Hallucination (rare <5%):",
        "   • Problem: Invented flight numbers not in context",
        "   • Solution: Stronger grounding in prompt + validation",
        "",
        "2. Incomplete Responses (Mistral):",
        "   • Problem: Single character/word responses",
        "   • Solution: Switch to Llama/Gemma, add validation",
        "",
        "3. Ignoring Constraints:",
        "   • Problem: Recommended class not in context",
        "   • Solution: Emphasize constraints, post-process filter"
    ]
)

# Section 5: Improvements
add_section_slide("5. Improvements Added")

add_content_slide(
    "Key Enhancements",
    [
        "✓ Hybrid Retrieval Strategy",
        "  • Combined symbolic + semantic",
        "  • Increased recall by 40%",
        "",
        "✓ Structured Prompt Engineering",
        "  • 3-part: Persona + Context + Task",
        "  • Reduced hallucination from 30% → <5%",
        "",
        "✓ Multiple Model Support",
        "  • Easy model switching",
        "  • A/B testing capability",
        "  • Fallback if one fails"
    ]
)

add_content_slide(
    "More Improvements",
    [
        "✓ Deduplication & Ranking",
        "  • Remove duplicate flights from both sources",
        "  • Smart ranking: delay → food → similarity",
        "",
        "✓ Context Limitation",
        "  • Top 15 results to stay within token budget",
        "  • Faster LLM responses",
        "",
        "✓ Comprehensive Error Handling",
        "  • No crashes from API failures",
        "  • Informative error messages",
        "",
        "✓ Model Comparison Framework",
        "  • Automated testing with metrics",
        "  • CSV export for analysis"
    ]
)

# Section 6: Limitations
add_section_slide("6. Remaining Limitations")

add_content_slide(
    "Data & Retrieval Limitations",
    [
        "Data:",
        "  • Limited attributes (no price, amenities, Wi-Fi)",
        "  • Small dataset (2,000 journeys)",
        "  • No real-time availability",
        "",
        "Retrieval:",
        "  • No personalization (same for all users)",
        "  • Limited multi-hop reasoning",
        "  • Static embeddings (TF-IDF/BoW)",
        "",
        "Future: Expand KG, use transformers, add personalization"
    ]
)

add_content_slide(
    "System Limitations",
    [
        "LLM:",
        "  • No real-time learning from feedback",
        "  • Prompt sensitivity",
        "  • Cost & latency (3-15s, $0.001-0.01/request)",
        "",
        "System:",
        "  • No conversational context (each query independent)",
        "  • No booking integration",
        "  • Limited error recovery",
        "",
        "Scalability:",
        "  • Single Neo4j instance",
        "  • No load balancing"
    ]
)

# Section 7: Conclusion
add_section_slide("7. Conclusion")

add_content_slide(
    "Key Achievements",
    [
        "✅ End-to-end RAG system (KG + LLM)",
        "✅ Hybrid retrieval for better recall",
        "✅ Multi-model comparison framework",
        "✅ Production-ready UI (Streamlit)",
        "✅ Comprehensive evaluation",
        "",
        "Best Practices:",
        "  • Structured prompting reduces hallucination",
        "  • Hybrid retrieval outperforms single approach",
        "  • Model choice depends on use case",
        "  • Error analysis guides improvements"
    ]
)

add_content_slide(
    "Future Roadmap",
    [
        "Short-term (1-3 months):",
        "  • Add pricing data to KG",
        "  • Conversation history",
        "  • Upgrade to Sentence-BERT embeddings",
        "",
        "Medium-term (3-6 months):",
        "  • Multi-hop query support",
        "  • Personalization engine",
        "  • Real-time flight API integration",
        "",
        "Long-term (6-12 months):",
        "  • Self-hosted LLM (cost reduction)",
        "  • Multi-language support",
        "  • Full booking integration"
    ]
)

add_content_slide(
    "Lessons Learned",
    [
        "1. Hybrid approaches work best",
        "   Combining symbolic + semantic > either alone",
        "",
        "2. Prompt engineering is critical",
        "   Small changes significantly impact quality",
        "",
        "3. Model selection matters",
        "   Quality vs speed vs cost trade-offs",
        "",
        "4. Error analysis is essential",
        "   Understanding failures guides improvements",
        "",
        "5. Comprehensive evaluation needed",
        "   Both quantitative AND qualitative metrics"
    ]
)

# Final slide
add_content_slide(
    "Thank You!",
    [
        "Questions?",
        "",
        "Project Repository:",
        "  github.com/yourusername/airline-assistant",
        "",
        "Documentation:",
        "  • FINAL_PRESENTATION_REPORT.md",
        "  • model_comparison_summary.csv",
        "  • MODEL_COMPARISON_GUIDE.md",
        "",
        "Demo: streamlit run ui_D/app.py"
    ]
)

# Save presentation
output_file = "Airline_Assistant_Presentation.pptx"
prs.save(output_file)

print(f"✅ PowerPoint presentation created: {output_file}")
print(f"   Total slides: {len(prs.slides)}")
print("\nYou can open this file in:")
print("  • Microsoft PowerPoint")
print("  • Google Slides (upload)")
print("  • LibreOffice Impress")
